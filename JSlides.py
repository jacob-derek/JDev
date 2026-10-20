#!/usr/bin/env python3
"""
JSlides - visualizador de apresentações para o JDesk (Fase 4).

    python JSlides.py arquivo.pptx|.ppt|.odp|.pdf

Como funciona
  * PDF:  aberto direto com PyMuPDF e desenhado no tamanho exato do painel (nítido).
  * PPTX/PPT/ODP: convertido para PDF pelo LibreOffice (headless, numa thread, com cache em
    ~/.cache/jdesk/slides por caminho+mtime+tamanho: a 2ª abertura é instantânea).
  * Sem LibreOffice ou sem PyMuPDF: modo texto (python-pptx) com títulos, corpo e notas.
  * Dentro do JDesk ele é um painel que cobre o workspace; Esc/q fecha e devolve o foco.

Teclas
    → ↓ PgDn espaço l j       próximo slide          ← ↑ PgUp Backspace h k   slide anterior
    Home / End                primeiro / último      g <número> Enter          ir para o slide
    o                         grade de miniaturas    n                         notas do orador
    f                         tela limpa             r                         recarregar
    q / Esc                   fechar

Dependências: pip install pymupdf python-pptx   +   LibreOffice (soffice) para .pptx/.ppt/.odp
"""
from __future__ import annotations

import hashlib
import os
import shutil
import subprocess
import sys
import threading
import time
import tkinter as tk
from pathlib import Path
from tkinter import font as tkfont

try:
    import pymupdf as fitz
except ImportError:
    try:
        import fitz  # nome antigo do mesmo pacote
    except ImportError:
        fitz = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

CACHE_DIR = Path(os.environ.get("XDG_CACHE_HOME") or Path.home() / ".cache") / "jdesk" / "slides"
CACHE_MAX = 40  # PDFs convertidos mantidos em cache
EXT_PDF = {".pdf"}
EXT_OFFICE = {".pptx", ".ppt", ".odp", ".pps", ".ppsx", ".potx"}

# Tokyo Night
BG, PAINEL, FG = "#1a1b26", "#16161e", "#c0caf5"
DIM, AZUL, AMARELO, VERMELHO = "#565f89", "#7aa2f7", "#e0af68", "#f7768e"


class ErroSlides(Exception):
    """Erro esperado, com mensagem pronta para o usuário."""


# ══════════════════════════════════════════════════════════════════════════════
#  Lógica sem interface (testável)
# ══════════════════════════════════════════════════════════════════════════════
def chave_cache(caminho: Path) -> str:
    st = caminho.stat()
    bruto = f"{caminho.resolve()}|{st.st_mtime_ns}|{st.st_size}"
    return hashlib.sha1(bruto.encode("utf-8")).hexdigest()[:20]


def pdf_em_cache(caminho: Path) -> Path:
    return CACHE_DIR / f"{chave_cache(caminho)}.pdf"


def achar_soffice() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def podar_cache() -> None:
    """Mantém só os CACHE_MAX PDFs mais recentes."""
    try:
        pdfs = sorted(CACHE_DIR.glob("*.pdf"), key=lambda p: p.stat().st_mtime, reverse=True)
        for velho in pdfs[CACHE_MAX:]:
            velho.unlink(missing_ok=True)
    except OSError:
        pass


def converter_para_pdf(origem: Path, destino: Path, timeout: int = 180) -> None:
    """Converte com o LibreOffice para `destino` (escrita atômica). Levanta ErroSlides."""
    soffice = achar_soffice()
    if not soffice:
        raise ErroSlides("LibreOffice (soffice) não encontrado.")
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    saida = CACHE_DIR / f"tmp-{os.getpid()}-{threading.get_ident()}"
    perfil = CACHE_DIR / "lo-profile"  # perfil isolado: não colide com um LibreOffice aberto
    saida.mkdir(parents=True, exist_ok=True)
    try:
        cmd = [soffice, f"-env:UserInstallation={perfil.as_uri()}", "--headless",
               "--convert-to", "pdf", "--outdir", str(saida), str(origem)]
        try:
            r = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout,
                               stdin=subprocess.DEVNULL)
        except subprocess.TimeoutExpired:
            raise ErroSlides(f"A conversão passou de {timeout}s e foi cancelada.") from None
        gerado = saida / f"{origem.stem}.pdf"
        if not gerado.exists():
            detalhe = (r.stderr or r.stdout or "").strip().splitlines()
            raise ErroSlides("O LibreOffice não gerou o PDF." + (f" ({detalhe[-1]})" if detalhe else ""))
        os.replace(gerado, destino)
    finally:
        shutil.rmtree(saida, ignore_errors=True)
    podar_cache()


def _texto_forma(forma) -> list[str]:
    linhas: list[str] = []
    if getattr(forma, "has_text_frame", False) and forma.text_frame is not None:
        for par in forma.text_frame.paragraphs:
            t = "".join(r.text for r in par.runs).strip() or par.text.strip()
            if t:
                linhas.append(("  " * par.level) + t)
    if getattr(forma, "has_table", False) and forma.has_table:
        for linha in forma.table.rows:
            linhas.append(" | ".join(c.text.strip() for c in linha.cells))
    if getattr(forma, "shapes", None) is not None and forma.shape_type == 6:  # grupo
        for sub in forma.shapes:
            linhas.extend(_texto_forma(sub))
    return linhas


def ler_pptx(caminho: Path) -> list[dict]:
    """[{titulo, corpo: [linhas], notas}] por slide, via python-pptx."""
    if Presentation is None:
        raise ErroSlides("python-pptx não instalado (pip install python-pptx).")
    try:
        prs = Presentation(str(caminho))
    except Exception as exc:
        raise ErroSlides(f"Não consegui ler o pptx: {exc}") from exc
    slides = []
    for i, s in enumerate(prs.slides, 1):
        titulo = ""
        if s.shapes.title is not None and s.shapes.title.has_text_frame:
            titulo = s.shapes.title.text_frame.text.strip()
        corpo: list[str] = []
        for forma in s.shapes:
            if s.shapes.title is not None and forma.shape_id == s.shapes.title.shape_id:
                continue
            corpo.extend(_texto_forma(forma))
        notas = ""
        if s.has_notes_slide and s.notes_slide.notes_text_frame is not None:
            notas = s.notes_slide.notes_text_frame.text.strip()
        slides.append({"titulo": titulo or f"Slide {i}", "corpo": corpo, "notas": notas})
    return slides


# ══════════════════════════════════════════════════════════════════════════════
#  Painel
# ══════════════════════════════════════════════════════════════════════════════
class SlidesPane(tk.Frame):
    es_viewer = True  # o JDesk usa isto para tratar teclas de forma diferente dos painéis pty

    def __init__(self, parent, on_close=None, key_filter=None, on_fullscreen=None):
        super().__init__(parent, bg=BG, bd=0, highlightthickness=0)
        self.on_close = on_close
        self.key_filter = key_filter
        self.on_fullscreen = on_fullscreen

        base = tkfont.nametofont("TkDefaultFont")
        self.f_status = base.copy()
        self.f_status.configure(size=10)
        self.f_titulo = base.copy()
        self.f_titulo.configure(size=22, weight="bold")
        self.f_corpo = base.copy()
        self.f_corpo.configure(size=14)
        self.f_msg = base.copy()
        self.f_msg.configure(size=13)

        self.status = tk.Label(self, bg=PAINEL, fg=DIM, anchor="w", font=self.f_status, padx=8, pady=3)
        self.status.pack(side=tk.BOTTOM, fill=tk.X)
        self.notas = tk.Text(self, height=7, bg=PAINEL, fg=FG, wrap="word", bd=0, padx=12, pady=8,
                             highlightthickness=0, font=self.f_corpo, state="disabled")
        self.canvas = tk.Canvas(self, bg=BG, highlightthickness=0, bd=0, takefocus=1)
        self.canvas.pack(fill=tk.BOTH, expand=True)

        self.doc = None            # documento PyMuPDF
        self.texto: list[dict] | None = None   # modo texto (python-pptx)
        self.notas_lista: list[str] = []
        self.caminho: Path | None = None
        self.titulo = ""
        self.total = 0
        self.pagina = 0
        self.modo = "slide"        # "slide" | "grade"
        self.msg = ""              # mensagem central (convertendo..., erro)
        self.aviso = ""            # aviso curto na barra de status
        self.goto_buf: str | None = None
        self.mostrar_notas = False
        self.tela_limpa = False
        self.grade_sel = 0
        self._img = None
        self._cache_img: dict = {}
        self._grade_imgs: dict = {}
        self._geracao = 0
        self._destruido = False
        self._redesenho = None

        self.canvas.bind("<Key>", self.on_key)
        self.canvas.bind("<Button-1>", lambda e: self.canvas.focus_set())
        self.canvas.bind("<Configure>", self._ao_redimensionar)
        self.canvas.bind("<MouseWheel>", lambda e: self.ir(-1 if e.delta > 0 else 1))
        self.canvas.bind("<Button-4>", lambda e: self.ir(-1))
        self.canvas.bind("<Button-5>", lambda e: self.ir(1))

    # ── abertura ──────────────────────────────────────────────────────────
    def focus_canvas(self):
        self.canvas.focus_set()

    def abrir(self, caminho) -> None:
        self._limpar()
        p = Path(caminho).expanduser()
        self.caminho, self.titulo = p, p.name
        if not p.is_file():
            return self._erro(f"Arquivo não encontrado: {p}")
        ext = p.suffix.lower()
        if ext in EXT_PDF:
            self._carregar_pdf(p)
        elif ext in EXT_OFFICE:
            self._abrir_office(p)
        else:
            self._erro(f"Formato não suportado: {ext or '(sem extensão)'}")

    def _limpar(self):
        self._geracao += 1  # invalida conversões em andamento
        if self.doc is not None:
            try:
                self.doc.close()
            except Exception:
                pass
        self.doc, self.texto, self.notas_lista = None, None, []
        self.total = self.pagina = self.grade_sel = 0
        self.modo, self.msg, self.aviso, self.goto_buf = "slide", "", "", None
        self._cache_img.clear()
        self._grade_imgs.clear()
        self._img = None

    def _erro(self, texto: str):
        self.msg = texto
        self.desenhar()

    def _abrir_office(self, p: Path):
        pdf = pdf_em_cache(p)
        if fitz and pdf.exists():
            return self._carregar_pdf(pdf, origem=p)
        if fitz and achar_soffice():
            return self._converter_async(p, pdf)
        motivo = ("PyMuPDF não instalado (pip install pymupdf)" if not fitz
                  else "LibreOffice não encontrado")
        self._carregar_texto(p, f"{motivo}: modo texto")

    def _converter_async(self, origem: Path, pdf: Path):
        gen = self._geracao
        resultado: dict = {}

        def trabalho():
            try:
                converter_para_pdf(origem, pdf)
                resultado["ok"] = True
            except ErroSlides as exc:
                resultado["erro"] = str(exc)
            except Exception as exc:  # noqa: BLE001
                resultado["erro"] = repr(exc)

        t = threading.Thread(target=trabalho, daemon=True)
        t.start()
        inicio = time.monotonic()

        def vigiar():
            if self._destruido or gen != self._geracao:
                return
            if t.is_alive():
                self.msg = f"Convertendo para PDF (LibreOffice)…  {int(time.monotonic() - inicio)}s"
                self.desenhar()
                self.after(400, vigiar)
                return
            if resultado.get("ok"):
                self.msg = ""
                self._carregar_pdf(pdf, origem=origem)
            else:
                self._carregar_texto(origem, f"Falha na conversão ({resultado.get('erro')}): modo texto")

        self.msg = "Convertendo para PDF (LibreOffice)…"
        self.desenhar()
        self.after(250, vigiar)

    def _carregar_pdf(self, pdf: Path, origem: Path | None = None):
        if not fitz:
            return self._erro("PyMuPDF não instalado: pip install pymupdf")
        try:
            self.doc = fitz.open(str(pdf))
        except Exception as exc:  # noqa: BLE001
            return self._erro(f"Não consegui abrir o PDF: {exc}")
        self.total, self.pagina, self.modo, self.msg = len(self.doc), 0, "slide", ""
        fonte = origem or pdf
        if fonte.suffix.lower() == ".pptx" and Presentation is not None:
            try:
                self.notas_lista = [s["notas"] for s in ler_pptx(fonte)]
            except ErroSlides:
                self.notas_lista = []
        self.desenhar()

    def _carregar_texto(self, p: Path, aviso: str):
        try:
            if p.suffix.lower() != ".pptx":
                raise ErroSlides("Sem LibreOffice/PyMuPDF só consigo mostrar .pptx em modo texto.")
            self.texto = ler_pptx(p)
        except ErroSlides as exc:
            return self._erro(str(exc))
        self.total, self.pagina, self.msg, self.aviso = len(self.texto), 0, "", aviso
        self.notas_lista = [s["notas"] for s in self.texto]
        self.desenhar()

    # ── navegação ─────────────────────────────────────────────────────────
    def ir(self, delta: int):
        if self.total:
            self.ver(self.pagina + delta)

    def ver(self, n: int):
        if not self.total:
            return
        n = max(0, min(self.total - 1, n))
        if n != self.pagina:
            self.pagina = n
            self.desenhar()

    def fechar(self):
        self._limpar()
        if self.on_close:
            self.on_close()

    def destroy(self):
        self._destruido = True
        self._limpar()
        super().destroy()

    # ── teclado ───────────────────────────────────────────────────────────
    def on_key(self, e):
        if self.key_filter and self.key_filter(self, e):
            return "break"
        ks = e.keysym
        if ks in ("Shift_L", "Shift_R", "Control_L", "Control_R", "Alt_L", "Alt_R", "Meta_L", "Meta_R",
                  "Super_L", "Super_R", "Caps_Lock", "ISO_Level3_Shift"):
            return "break"

        if self.goto_buf is not None:  # modo "ir para o slide N"
            if ks in ("Return", "KP_Enter"):
                if self.goto_buf:
                    self.ver(int(self.goto_buf) - 1)
                self.goto_buf = None
            elif ks == "Escape":
                self.goto_buf = None
            elif ks == "BackSpace":
                self.goto_buf = self.goto_buf[:-1]
            elif e.char.isdigit() and len(self.goto_buf) < 5:
                self.goto_buf += e.char
            self._status()
            return "break"

        if self.modo == "grade":
            return self._tecla_grade(ks)

        if ks in ("Right", "Down", "Next", "space", "l", "j"):
            self.ir(1)
        elif ks in ("Left", "Up", "Prior", "BackSpace", "h", "k"):
            self.ir(-1)
        elif ks == "Home":
            self.ver(0)
        elif ks == "End":
            self.ver(self.total - 1)
        elif ks == "g" and self.total:
            self.goto_buf = ""
            self._status()
        elif ks == "o":
            self._abrir_grade()
        elif ks == "n":
            self._alternar_notas()
        elif ks == "f":
            self._tela_limpa()
        elif ks == "r" and self.caminho:
            pagina = self.pagina
            self.abrir(self.caminho)
            if self.total:
                self.ver(pagina)
        elif ks in ("q", "Escape"):
            self.fechar()
        return "break"

    def _tecla_grade(self, ks):
        cols = self._grade_geometria()[0]
        if ks in ("Right", "l"):
            self.grade_sel = min(self.total - 1, self.grade_sel + 1)
        elif ks in ("Left", "h"):
            self.grade_sel = max(0, self.grade_sel - 1)
        elif ks in ("Down", "j"):
            self.grade_sel = min(self.total - 1, self.grade_sel + cols)
        elif ks in ("Up", "k"):
            self.grade_sel = max(0, self.grade_sel - cols)
        elif ks == "Home":
            self.grade_sel = 0
        elif ks == "End":
            self.grade_sel = self.total - 1
        elif ks in ("Return", "KP_Enter", "space"):
            self.pagina, self.modo = self.grade_sel, "slide"
        elif ks in ("Escape", "o", "q"):
            self.modo = "slide"
        self.desenhar()
        return "break"

    def _alternar_notas(self):
        self.mostrar_notas = not self.mostrar_notas
        if self.mostrar_notas:
            self.notas.pack(side=tk.BOTTOM, fill=tk.X, before=self.canvas)
        else:
            self.notas.pack_forget()
        self.desenhar()

    def _tela_limpa(self):
        self.tela_limpa = not self.tela_limpa
        if self.tela_limpa:
            self.status.pack_forget()
        else:
            self.status.pack(side=tk.BOTTOM, fill=tk.X, before=self.canvas)
        if self.on_fullscreen:
            self.on_fullscreen(self.tela_limpa)
        self.desenhar()

    # ── desenho ───────────────────────────────────────────────────────────
    def _ao_redimensionar(self, _e):
        if self._redesenho:
            self.after_cancel(self._redesenho)
        self._redesenho = self.after(80, self.desenhar)

    def desenhar(self):
        if self._destruido:
            return
        c = self.canvas
        w, h = c.winfo_width(), c.winfo_height()
        c.delete("all")
        if w < 40 or h < 40:
            return
        if self.msg and not self.total:
            cor = VERMELHO if self.msg.startswith(("Não", "Arquivo", "Formato", "PyMuPDF", "Sem ")) else DIM
            c.create_text(w // 2, h // 2, text=self.msg, fill=cor, font=self.f_msg, width=w - 80,
                          justify="center")
        elif not self.total:
            c.create_text(w // 2, h // 2, text="Nenhum arquivo aberto.", fill=DIM, font=self.f_msg)
        elif self.modo == "grade":
            self._desenhar_grade(w, h)
        elif self.doc is not None:
            self._desenhar_pdf(w, h)
        elif self.texto is not None:
            self._desenhar_texto(w, h)
        self._atualizar_notas()
        self._status()

    def _desenhar_pdf(self, w, h):
        chave = (self.pagina, w, h)
        img = self._cache_img.get(chave)
        if img is None:
            try:
                page = self.doc[self.pagina]
                r = page.rect
                esc = min((w - 8) / r.width, (h - 8) / r.height)
                pix = page.get_pixmap(matrix=fitz.Matrix(esc, esc), alpha=False)
                img = tk.PhotoImage(data=pix.tobytes("ppm"))
            except Exception as exc:  # noqa: BLE001
                self.canvas.create_text(w // 2, h // 2, text=f"Erro ao desenhar o slide: {exc}",
                                        fill=VERMELHO, font=self.f_msg, width=w - 80)
                return
            if len(self._cache_img) > 8:
                self._cache_img.clear()
            self._cache_img[chave] = img
        self._img = img
        self.canvas.create_image(w // 2, h // 2, image=img)

    def _desenhar_texto(self, w, h):
        s = self.texto[self.pagina]
        x, y, larg = 48, 40, max(100, w - 96)
        self.canvas.create_text(x, y, text=s["titulo"], anchor="nw", fill=AZUL, font=self.f_titulo, width=larg)
        corpo = "\n".join(s["corpo"]) or "(sem texto)"
        self.canvas.create_text(x, y + 70, text=corpo, anchor="nw", fill=FG, font=self.f_corpo, width=larg)

    def _atualizar_notas(self):
        if not self.mostrar_notas:
            return
        nota = self.notas_lista[self.pagina] if self.pagina < len(self.notas_lista) else ""
        self.notas.configure(state="normal")
        self.notas.delete("1.0", "end")
        self.notas.insert("1.0", nota or "(sem notas neste slide)")
        self.notas.configure(state="disabled")

    def _status(self):
        if self.tela_limpa:
            return
        partes = [self.titulo]
        if self.total:
            partes.append(f"{self.pagina + 1}/{self.total}")
        if self.goto_buf is not None:
            partes.append(f"ir para: {self.goto_buf}_")
        elif self.aviso:
            partes.append(self.aviso)
        else:
            partes.append("←→ navegar · g ir · o grade · n notas · f limpa · r recarrega · q fecha")
        self.status.configure(text="   ·   ".join(p for p in partes if p))

    # ── grade de miniaturas ───────────────────────────────────────────────
    TW, PAD = 220, 16

    def _abrir_grade(self):
        if self.doc is None:
            self.aviso = "Grade só funciona com PDF/conversão (modo texto não tem miniaturas)."
            self._status()
            return
        self.modo, self.grade_sel = "grade", self.pagina
        self.desenhar()

    def _grade_geometria(self):
        w = max(self.canvas.winfo_width(), 1)
        cols = max(1, (w - self.PAD) // (self.TW + self.PAD))
        if self.doc is not None and self.total:
            r = self.doc[0].rect
            th = int(self.TW * r.height / r.width)
        else:
            th = int(self.TW * 9 / 16)
        return cols, th

    def _miniatura(self, i, th):
        img = self._grade_imgs.get(i)
        if img is None:
            r = self.doc[i].rect
            esc = self.TW / r.width
            pix = self.doc[i].get_pixmap(matrix=fitz.Matrix(esc, esc), alpha=False)
            img = tk.PhotoImage(data=pix.tobytes("ppm"))
            self._grade_imgs[i] = img
        return img

    def _desenhar_grade(self, w, h):
        cols, th = self._grade_geometria()
        celula_h = th + self.PAD + 18
        linhas_vis = max(1, h // celula_h)
        sel_linha = self.grade_sel // cols
        topo = max(0, min(sel_linha - linhas_vis // 2, (self.total - 1) // cols - linhas_vis + 1))
        topo = max(0, topo)
        for idx in range(topo * cols, min(self.total, (topo + linhas_vis + 1) * cols)):
            lin, col = divmod(idx, cols)
            x = self.PAD + col * (self.TW + self.PAD)
            y = self.PAD + (lin - topo) * celula_h
            try:
                self.canvas.create_image(x, y, image=self._miniatura(idx, th), anchor="nw")
            except Exception:  # noqa: BLE001
                self.canvas.create_rectangle(x, y, x + self.TW, y + th, outline=DIM)
            sel = idx == self.grade_sel
            self.canvas.create_rectangle(x - 2, y - 2, x + self.TW + 2, y + th + 2,
                                         outline=AMARELO if sel else PAINEL, width=3 if sel else 1)
            self.canvas.create_text(x, y + th + 4, text=str(idx + 1), anchor="nw",
                                    fill=AMARELO if sel else DIM, font=self.f_status)


# ══════════════════════════════════════════════════════════════════════════════
#  Execução isolada
# ══════════════════════════════════════════════════════════════════════════════
def main() -> int:
    if len(sys.argv) < 2:
        print("Uso: python JSlides.py arquivo.pptx|.ppt|.odp|.pdf")
        return 1
    root = tk.Tk()
    root.title("JSlides")
    root.geometry("1100x700")
    root.configure(bg=BG)
    pane = SlidesPane(root, on_close=root.destroy,
                      on_fullscreen=lambda on: root.attributes("-fullscreen", on))
    pane.pack(fill=tk.BOTH, expand=True)

    def iniciar():
        pane.abrir(sys.argv[1])
        pane.focus_canvas()

    root.after(80, iniciar)
    root.mainloop()
    return 0


if __name__ == "__main__":
    sys.exit(main())
